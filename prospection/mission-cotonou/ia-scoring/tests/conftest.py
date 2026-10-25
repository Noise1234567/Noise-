"""Classeur de test et fausses réponses du modèle. Aucun test n'appelle Ollama."""

import json

import pytest
from openpyxl import Workbook

# Observations fictives, assez longues (plus de 200 caractères) pour une confiance normale.
LONG_TEXT = (
    "Affiche publiée le 3 octobre : « Soirée Neon Vibes, organisée par Collectif Exemple, "
    "4e édition, complet la semaine dernière ». Réservez vos places en message privé. "
    "Prochaine date annoncée pour novembre au même club de Cotonou, entrée 5 000 FCFA."
)

PROSPECTS_HEADERS = [
    "ID prospect", "Nom du compte ou de l'organisation", "Exemple fictif (oui/non)",
    "Observations publiques collées", "Statut analyse IA",
]
SCORES_HEADERS = [
    "ID prospect", "Compte", "Pc : probabilité de conversion (0 à 100)",
    "I : impact commercial (1 à 5)", "R : récurrence (1 à 5)", "N : recommandation (1 à 5)",
    "F : facilité de contact (1 à 5)", "Score final (/100)", "Priorité",
    "Confiance des données", "Justification du classement", "Inconnues à vérifier",
    "Pe : probabilité d'échec (0 à 100)", "G : gravité en cas d'échec (1 à 5)", "Persona",
    "Citation du persona", "Avantages", "Inconvénients", "Confiance IA", "Relecture",
]


def make_workbook(prospects, with_new_columns=True):
    """prospects : liste de tuples (id, exemple, observations, statut), à partir de la ligne 5."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Prospects"
    headers = PROSPECTS_HEADERS if with_new_columns else PROSPECTS_HEADERS[:3]
    for col, header in enumerate(headers, start=1):
        ws.cell(4, col, header)
    for row, (pid, example, text, status) in enumerate(prospects, start=5):
        ws.cell(row, 1, pid)
        ws.cell(row, 2, f"Compte {pid}")
        ws.cell(row, 3, example)
        if with_new_columns:
            ws.cell(row, 4, text)
            ws.cell(row, 5, status)
    scores = wb.create_sheet("Scores")
    headers = SCORES_HEADERS if with_new_columns else SCORES_HEADERS[:12]
    for col, header in enumerate(headers, start=1):
        scores.cell(4, col, header)
    return wb


def model_answer(**overrides):
    """Réponse valide du modèle pour LONG_TEXT ; chaque champ peut être remplacé."""
    answer = {
        "pc": 70, "i": 4, "r": 5, "n": 4, "f": 3, "pe": 40, "g": 4,
        "persona": "P5",
        "persona_citation": "organisée par Collectif Exemple",
        "avantages": [{"point": "Vend déjà par messages : besoin fort",
                       "citation": "Réservez vos places en message privé"}],
        "inconvenients": [],
        "confiance": "Moyenne",
        "justification": "Structure qui signe des éditions répétées.",
        "inconnues": "Nombre réel de participants",
    }
    answer.update(overrides)
    return json.dumps(answer, ensure_ascii=False)


class FakeModel:
    """Remplace Ollama : renvoie les réponses prévues, dans l'ordre, et garde les messages reçus."""

    def __init__(self, *answers):
        self.answers = list(answers)
        self.received = []

    def __call__(self, messages):
        self.received.append(messages)
        answer = self.answers.pop(0)
        if isinstance(answer, Exception):
            raise answer
        return answer


@pytest.fixture
def long_text():
    return LONG_TEXT
